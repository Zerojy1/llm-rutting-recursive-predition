from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree


SOURCE = Path(r"C:\Users\zianyy\Desktop\Qwen-generated state transitions for long-horizon pavement rutting prediction A TDR-PRSN study on RIOHTrack_restructured.docx")
OUTPUT = Path(r"C:\Users\zianyy\Desktop\Qwen-generated state transitions for long-horizon pavement rutting prediction A TDR-PRSN study on RIOHTrack_all_native_equations.docx")
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")


MATHML_ROLLOUT = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <msub>
      <mover accent="true"><mi>r</mi><mo>^</mo></mover>
      <mrow><mi>t</mi><mo>+</mo><mi>k</mi></mrow>
    </msub>
    <mo>=</mo>
    <msub><mi>F</mi><mi>θ</mi></msub>
    <mo>(</mo>
    <mi mathvariant="bold">s</mi><mo>,</mo>
    <msub>
      <mover accent="true"><mi mathvariant="bold">r</mi><mo>^</mo></mover>
      <mrow>
        <mi>t</mi><mo>+</mo><mi>k</mi><mo>−</mo><mn>4</mn><mo>:</mo>
        <mi>t</mi><mo>+</mo><mi>k</mi><mo>−</mo><mn>1</mn>
      </mrow>
    </msub>
    <mo>,</mo>
    <msub><mi mathvariant="bold">u</mi><mrow><mi>t</mi><mo>+</mo><mi>k</mi></mrow></msub>
    <mo>)</mo>
    <mo>,</mo><mspace width="1em"/>
    <msub><mover accent="true"><mi>r</mi><mo>^</mo></mover><mi>t</mi></msub>
    <mo>=</mo><msub><mi>r</mi><mi>t</mi></msub>
    <mo>,</mo><mspace width="1em"/>
    <mi>k</mi><mo>=</mo><mn>1</mn><mo>,</mo><mo>…</mo><mo>,</mo><mi>H</mi>
  </mrow>
</math>
"""


MATHML_GENERATOR = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <msub><mi>q</mi><mi>φ</mi></msub>
    <mo>(</mo>
    <msub>
      <mrow><mo>Δ</mo><mi mathvariant="bold">r</mi></mrow>
      <mrow><mi>t</mi><mo>+</mo><mn>1</mn><mo>:</mo><mi>t</mi><mo>+</mo><mn>6</mn></mrow>
    </msub>
    <mo>|</mo>
    <mi mathvariant="bold">s</mi><mo>,</mo>
    <msub>
      <mi mathvariant="bold">r</mi>
      <mrow><mi>t</mi><mo>−</mo><mn>4</mn><mo>:</mo><mi>t</mi></mrow>
    </msub>
    <mo>,</mo>
    <msub>
      <mi mathvariant="bold">u</mi>
      <mrow><mi>t</mi><mo>+</mo><mn>1</mn><mo>:</mo><mi>t</mi><mo>+</mo><mn>6</mn></mrow>
    </msub>
    <mo>)</mo>
    <mo>,</mo><mspace width="1em"/>
    <msub><mrow><mo>Δ</mo><mi>r</mi></mrow><mrow><mi>t</mi><mo>+</mo><mi>j</mi></mrow></msub>
    <mo>=</mo>
    <msub><mi>r</mi><mrow><mi>t</mi><mo>+</mo><mi>j</mi></mrow></msub>
    <mo>−</mo>
    <msub><mi>r</mi><mrow><mi>t</mi><mo>+</mo><mi>j</mi><mo>−</mo><mn>1</mn></mrow></msub>
  </mrow>
</math>
"""


MATHML_RECONSTRUCTION = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <msub><mover accent="true"><mi>r</mi><mo>~</mo></mover><mrow><mi>t</mi><mo>+</mo><mi>j</mi></mrow></msub>
    <mo>=</mo><msub><mi>r</mi><mi>t</mi></msub><mo>+</mo>
    <munderover>
      <mo>∑</mo>
      <mrow><mi>k</mi><mo>=</mo><mn>1</mn></mrow>
      <mi>j</mi>
    </munderover>
    <msub>
      <mover accent="true"><mrow><mo>Δ</mo><mi>r</mi></mrow><mo>~</mo></mover>
      <mrow><mi>t</mi><mo>+</mo><mi>k</mi></mrow>
    </msub>
    <mo>,</mo><mspace width="1em"/>
    <mi>j</mi><mo>=</mo><mn>1</mn><mo>,</mo><mo>…</mo><mo>,</mo><mn>6</mn>
  </mrow>
</math>
"""


MATHML_IMPROVEMENT = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <msub><mi mathvariant="normal">Improvement</mi><mi>h</mi></msub><mo>(</mo><mo>%</mo><mo>)</mo>
    <mo>=</mo><mn>100</mn><mo>×</mo>
    <mfrac>
      <mrow>
        <msubsup><mi mathvariant="normal">RMSE</mi><mi>h</mi><mi mathvariant="normal">base</mi></msubsup>
        <mo>−</mo>
        <msubsup><mi mathvariant="normal">RMSE</mi><mi>h</mi><mi mathvariant="normal">aug</mi></msubsup>
      </mrow>
      <msubsup><mi mathvariant="normal">RMSE</mi><mi>h</mi><mi mathvariant="normal">base</mi></msubsup>
    </mfrac>
  </mrow>
</math>
"""


MATHML_SKILL = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <msub><mi mathvariant="normal">Skill</mi><mi>h</mi></msub>
    <mo>=</mo><mn>1</mn><mo>−</mo>
    <mfrac>
      <msub><mi mathvariant="normal">RMSE</mi><mrow><mi mathvariant="normal">model</mi><mo>,</mo><mi>h</mi></mrow></msub>
      <msub><mi mathvariant="normal">RMSE</mi><mrow><mi mathvariant="normal">persistence</mi><mo>,</mo><mi>h</mi></mrow></msub>
    </mfrac>
  </mrow>
</math>
"""


MATHML_MSSR = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mtable columnalign="left">
    <mtr><mtd><mrow>
      <msub><mi>d</mi><mi>h</mi></msub><mo>(</mo><mi>q</mi><mo>)</mo><mo>=</mo>
      <mfrac>
        <mrow>
          <msub><mi mathvariant="normal">RMSE</mi><mi>h</mi></msub><mo>(</mo><mi>q</mi><mo>)</mo>
          <mo>−</mo>
          <msub><mi mathvariant="normal">RMSE</mi><mi>h</mi></msub><mo>(</mo><mn>0</mn><mo>)</mo>
        </mrow>
        <mrow><msub><mi mathvariant="normal">RMSE</mi><mi>h</mi></msub><mo>(</mo><mn>0</mn><mo>)</mo></mrow>
      </mfrac>
    </mrow></mtd></mtr>
    <mtr><mtd><mrow>
      <msub><mi mathvariant="normal">MSSR</mi><mi>h</mi></msub><mo>(</mo><mi>τ</mi><mo>)</mo><mo>=</mo>
      <mi mathvariant="normal">max</mi><mo>{</mo><mn>0</mn><mo>}</mo><mo>∪</mo><mo>{</mo>
      <msub><mi>q</mi><mi>k</mi></msub><mo>∈</mo><mi>Q</mi><mo>:</mo>
      <msub><mi>U</mi><mi>h</mi></msub><mo>(</mo><msub><mi>q</mi><mi>j</mi></msub><mo>)</mo><mo>&lt;</mo><mi>τ</mi><mo>,</mo>
      <mo>∀</mo><mi>j</mi><mo>≤</mo><mi>k</mi><mo>}</mo>
    </mrow></mtd></mtr>
    <mtr><mtd><mrow>
      <msub><mi mathvariant="normal">MSSR</mi><mi mathvariant="normal">all</mi></msub><mo>(</mo><mi>τ</mi><mo>)</mo><mo>=</mo>
      <munder><mi mathvariant="normal">min</mi><mrow><mi>h</mi><mo>∈</mo><mo>{</mo><mn>1</mn><mo>,</mo><mn>3</mn><mo>,</mo><mn>6</mn><mo>,</mo><mn>12</mn><mo>}</mo></mrow></munder>
      <msub><mi mathvariant="normal">MSSR</mi><mi>h</mi></msub><mo>(</mo><mi>τ</mi><mo>)</mo>
    </mrow></mtd></mtr>
  </mtable>
</math>
"""


def mathml_to_omml(mathml: str, transformer: etree.XSLT):
    source = etree.fromstring(mathml.encode("utf-8"))
    transformed = transformer(source)
    root = transformed.getroot()
    omath = root.find(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
    return omath if omath is not None else root


def clear_paragraph_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def append_tab(paragraph) -> None:
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:tab"))
    paragraph._p.append(run)


def append_text(paragraph, text: str) -> None:
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    paragraph._p.append(run)


def insert_numbered_equation(paragraph, omath, number: int) -> None:
    clear_paragraph_content(paragraph)
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("w:shd", "w:pBdr"):
        decoration = p_pr.find(qn(tag))
        if decoration is not None:
            p_pr.remove(decoration)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    stops = paragraph.paragraph_format.tab_stops
    stops.clear_all()
    stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    stops.add_tab_stop(Inches(6.40), WD_TAB_ALIGNMENT.RIGHT)
    append_tab(paragraph)
    paragraph._p.append(deepcopy(omath))
    append_tab(paragraph)
    append_text(paragraph, f"({number})")


doc = Document(SOURCE)
transformer = etree.XSLT(etree.parse(str(MML2OMML)))

targets = [
    ("r̂ₜ₊ₖ", MATHML_ROLLOUT, 1),
    ("q(Δrₜ₊₁", MATHML_GENERATOR, 2),
    ("[EQUATION 3 PLACEHOLDER]", MATHML_RECONSTRUCTION, 3),
    ("[EQUATION 4 PLACEHOLDER]", MATHML_IMPROVEMENT, 4),
    ("[EQUATION 5 PLACEHOLDER]", MATHML_SKILL, 5),
    ("d_h(q) = [RMSE_h(q)", MATHML_MSSR, 6),
]

for marker, mathml, number in targets:
    paragraph = next(p for p in doc.paragraphs if marker in p.text)
    omml = mathml_to_omml(mathml, transformer)
    insert_numbered_equation(paragraph, omml, number)

doc.save(OUTPUT)

# Verify that all equations are native Office Math objects, not text or images.
check = Document(OUTPUT)
omath_count = len(check.element.body.findall(".//" + qn("m:oMath")))
assert omath_count == 6, f"Expected six native equations, found {omath_count}"
assert not any("[EQUATION" in p.text for p in check.paragraphs)
print(OUTPUT)
print(f"native_omath_objects={omath_count}")
