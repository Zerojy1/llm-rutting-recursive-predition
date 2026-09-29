from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\zianyy\Desktop\Qwen-generated state transitions for long-horizon pavement rutting prediction A TDR-PRSN study on RIOHTrack_Figure1.docx")
IMAGE = Path(r"C:\Users\zianyy\Desktop\final\project\figures\reference_style_charts_v21\png\figure01_study_framework_simplified_cropped.png")
OUTPUT = Path(r"C:\Users\zianyy\Desktop\Qwen-generated state transitions for long-horizon pavement rutting prediction A TDR-PRSN study on RIOHTrack_restructured.docx")


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def replace_text_nodes(document: Document, old: str, new: str) -> None:
    for node in document.element.body.iter(qn("w:t")):
        if node.text and old in node.text:
            node.text = node.text.replace(old, new)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)


def set_three_line_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "bottom", "left", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single" if edge_name in ("top", "bottom") else "nil")
        if edge_name in ("top", "bottom"):
            edge.set(qn("w:sz"), "10")
            edge.set(qn("w:color"), "000000")

    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.find(qn("w:tcBorders"))
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = tc_borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            tc_borders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")


doc = Document(SOURCE)

# Replace the dense Figure 1 image and its obsolete multi-panel caption.
caption_index = next(
    i
    for i, p in enumerate(doc.paragraphs)
    if p.text.strip().startswith("Figure 1. Study design and experimental workflow.")
)
figure_paragraph = doc.paragraphs[caption_index - 1]
clear_paragraph(figure_paragraph)
figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure_paragraph.paragraph_format.space_before = Pt(6)
figure_paragraph.paragraph_format.space_after = Pt(6)
figure_paragraph.add_run().add_picture(str(IMAGE), width=Inches(6.5))

caption = doc.paragraphs[caption_index]
clear_paragraph(caption)
caption.add_run(
    "Figure 1. Overview of the generator–predictor framework. "
    "RIOHTrack observations provide both real updates and the source data for Qwen domain adaptation. "
    "Quality-controlled synthetic transitions are combined with real updates for recursive evaluation "
    "using TDR-PRSN and the GRU control."
)

# Make room for a dedicated chronology table and shift later table numbers.
replace_text_nodes(doc, "Table 4", "Table 5")
replace_text_nodes(doc, "Table 3", "Table 4")
replace_text_nodes(doc, "Table 2", "Table 3")

heading = next(p for p in doc.paragraphs if p.text.strip() == "2.2 Data split")
table_caption = insert_paragraph_after(heading, "Table 2. Chronological data allocation.")
table_caption.runs[0].bold = True
table_caption.runs[0].font.size = Pt(9)
table_caption.paragraph_format.keep_with_next = True

chronology = doc.add_table(rows=3, cols=4)
if doc.tables and len(doc.tables) > 1:
    source_tbl_pr = deepcopy(doc.tables[0]._tbl.tblPr)
    chronology._tbl.remove(chronology._tbl.tblPr)
    chronology._tbl.insert(0, source_tbl_pr)

rows = [
    ("Component", "Training / final fit", "Development", "Time-held-out use"),
    (
        "Qwen generator",
        "Train: 2016–2017; final refit: 2016–2018",
        "2018",
        "2019 transition generation",
    ),
    (
        "Downstream predictor",
        "2016–2019",
        "2020",
        "2021 chronological test",
    ),
]
for row_idx, values in enumerate(rows):
    for col_idx, value in enumerate(values):
        set_cell_text(chronology.cell(row_idx, col_idx), value, bold=row_idx == 0)
set_three_line_borders(chronology)

# Move the newly created table from the end of the document to immediately after its caption.
table_caption._p.addnext(chronology._tbl)

# Replace the two equation placeholders with concise displayed mathematical definitions.
equations = {
    "[EQUATION 1 PLACEHOLDER]": "r̂ₜ₊ₖ = F(s, r̂ₜ₊ₖ₋₄:ₜ₊ₖ₋₁, uₜ₊ₖ; θ),    r̂ₜ = rₜ,    k = 1, …, H",
    "[EQUATION 2 PLACEHOLDER]": "q(Δrₜ₊₁:ₜ₊₆ | s, rₜ₋₄:ₜ, uₜ₊₁:ₜ₊₆; φ),    Δrₜ₊ⱼ = rₜ₊ⱼ − rₜ₊ⱼ₋₁",
}
for marker, formula in equations.items():
    paragraph = next(p for p in doc.paragraphs if marker in p.text)
    clear_paragraph(paragraph)
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("w:shd", "w:pBdr"):
        decoration = p_pr.find(qn(tag))
        if decoration is not None:
            p_pr.remove(decoration)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(formula)
    run.font.name = "Cambria Math"
    run.font.size = Pt(10.5)

doc.save(OUTPUT)

# Structural validation after save.
check = Document(OUTPUT)
body_text = "\n".join(p.text for p in check.paragraphs)
assert "[EQUATION 1 PLACEHOLDER]" not in body_text
assert "[EQUATION 2 PLACEHOLDER]" not in body_text
assert "[FIGURE 1 PLACEHOLDER]" not in body_text
assert "Table 2. Chronological data allocation." in body_text
assert len(check.tables) == 5
print(OUTPUT)
print(f"tables={len(check.tables)}, inline_shapes={len(check.inline_shapes)}")
